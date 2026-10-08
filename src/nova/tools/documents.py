"""Real documents (PDF, Word, HTML, Markdown, text, CSV) delivered to the user as downloadable files."""

import io
import xml.etree.ElementTree as ElementTree
from pathlib import Path

import markdown

from nova.core.messages import Media
from nova.tools.base import Tool, ToolError, ToolOutput, object_schema
from nova.tools.creative import slugify
from nova.tools.file_access import FileAccessGuard

DEFAULT_OUTPUT_FOLDER = "NOVA"
FORMATS = ("pdf", "docx", "html", "md", "txt", "csv")
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
MARKDOWN_EXTENSIONS = ["tables", "fenced_code", "sane_lists"]

DOCUMENT_CSS = """
@page { size: a4 portrait; margin: 2cm; }
body { font-family: Helvetica, Arial, sans-serif; font-size: 11pt; line-height: 1.45; color: #1f2430; }
h1 { font-size: 22pt; color: #1a3a5c; margin: 0 0 12pt; }
h2 { font-size: 15pt; color: #2c5aa0; margin: 16pt 0 6pt; border-bottom: 1px solid #c9d6ea; padding-bottom: 3pt; }
h3 { font-size: 12pt; color: #1a3a5c; margin: 12pt 0 4pt; }
p { margin: 0 0 7pt; text-align: justify; }
ul, ol { margin: 0 0 8pt 16pt; }
table { width: 100%; margin: 8pt 0; }
th { background: #2c5aa0; color: #ffffff; font-weight: bold; }
th, td { border: 1px solid #9aaccc; padding: 4pt 6pt; text-align: left; }
code, pre { font-family: Courier, monospace; background: #f1f4f9; }
pre { padding: 6pt; }
blockquote { margin: 6pt 0 6pt 12pt; padding-left: 8pt; border-left: 3px solid #c9d6ea; color: #4a5468; }
"""


def markdown_to_html(content: str) -> str:
    return markdown.markdown(content, extensions=MARKDOWN_EXTENSIONS, output_format="xhtml")


def html_page(title: str, body_html: str) -> str:
    return (
        f'<!DOCTYPE html>\n<html lang="fr"><head><meta charset="utf-8"><title>{title}</title>'
        f"<style>{DOCUMENT_CSS}</style></head><body>{body_html}</body></html>"
    )


def render_pdf(title: str, content: str) -> bytes:
    from xhtml2pdf import pisa

    output = io.BytesIO()
    status = pisa.CreatePDF(html_page(title, markdown_to_html(content)), dest=output, encoding="utf-8")
    if status.err:
        raise ToolError("The PDF could not be generated from this content.")
    return output.getvalue()


def render_docx(content: str) -> bytes:
    import docx

    document = docx.Document()
    root = ElementTree.fromstring(f"<root>{markdown_to_html(content).replace('&nbsp;', '&#160;')}</root>")
    for element in root:
        add_docx_block(document, element)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def add_docx_block(document, element: ElementTree.Element) -> None:
    tag = element.tag
    if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        add_docx_runs(document.add_heading(level=int(tag[1])), element)
    elif tag in {"ul", "ol"}:
        style = "List Bullet" if tag == "ul" else "List Number"
        for item in element.findall("li"):
            add_docx_runs(document.add_paragraph(style=style), item)
    elif tag == "table":
        rows = element.findall(".//tr")
        if not rows:
            return
        columns = max(len(row) for row in rows)
        table = document.add_table(rows=len(rows), cols=columns)
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column_index, cell in enumerate(row):
                paragraph = table.cell(row_index, column_index).paragraphs[0]
                add_docx_runs(paragraph, cell, bold=cell.tag == "th")
    elif tag == "pre":
        document.add_paragraph("".join(element.itertext()), style="No Spacing")
    elif tag == "hr":
        document.add_paragraph("")
    else:
        add_docx_runs(document.add_paragraph(), element)


def add_docx_runs(paragraph, element: ElementTree.Element, bold: bool = False, italic: bool = False) -> None:
    if element.text:
        run = paragraph.add_run(element.text)
        run.bold, run.italic = bold or None, italic or None
    for child in element:
        if child.tag == "br":
            paragraph.add_run().add_break()
        elif child.tag in {"ul", "ol"}:
            paragraph.add_run("\n" + "\n".join("• " + "".join(item.itertext()) for item in child.findall("li")))
        else:
            add_docx_runs(
                paragraph,
                child,
                bold=bold or child.tag in {"strong", "b", "th"},
                italic=italic or child.tag in {"em", "i"},
            )
        if child.tail:
            run = paragraph.add_run(child.tail)
            run.bold, run.italic = bold or None, italic or None


def media_for(path: Path, title: str) -> Media:
    suffix = path.suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        kind = "image"
    elif suffix == ".svg":
        kind = "svg"
    elif suffix in {".html", ".htm"}:
        kind = "html"
    else:
        kind = "file"
    return Media(path=str(path), kind=kind, title=title)


class DocumentStudio:
    def __init__(self, access: FileAccessGuard) -> None:
        self.access = access

    def create_document(self, title: str, format: str, content: str, path: str | None = None) -> ToolOutput:
        if format not in FORMATS:
            raise ToolError(f"format must be one of: {', '.join(FORMATS)}.")
        destination = self._destination(title, format, path)
        data = self._render(title, format, content)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(data)
        return ToolOutput(
            text=(
                f"Document created: {destination}. The user now sees it in the chat with a preview, "
                "an Open button and a Download button; do not describe how to find it."
            ),
            media=[media_for(destination, title)],
        )

    def create_document_confirmation(self, title: str, format: str, content: str, path: str | None = None) -> str | None:
        if path is None:
            return None
        destination = self._destination(title, format, path)
        return f"Écraser le fichier existant {destination} ?" if destination.is_file() else None

    def share_file(self, path: str, title: str | None = None) -> ToolOutput:
        file_path = self.access.resolve_for_read(path)
        if not file_path.is_file():
            raise ToolError(f"File not found: {file_path}")
        return ToolOutput(
            text=f"{file_path} is now shown to the user with Open and Download buttons.",
            media=[media_for(file_path, title or file_path.name)],
        )

    def _render(self, title: str, format: str, content: str) -> bytes:
        if format == "pdf":
            return render_pdf(title, content)
        if format == "docx":
            return render_docx(content)
        if format == "html":
            return html_page(title, markdown_to_html(content)).encode("utf-8")
        return content.encode("utf-8")

    def _destination(self, title: str, format: str, path: str | None) -> Path:
        if path is not None:
            destination = self.access.resolve_for_write(path)
            # The extension always matches the format: a "document" must never become a .bat or .exe.
            return destination if destination.suffix.lower() == f".{format}" else destination.with_name(f"{destination.name}.{format}")
        folder = self.access.resolve_for_write(DEFAULT_OUTPUT_FOLDER)
        candidate = folder / f"{slugify(title)}.{format}"
        counter = 2
        while candidate.exists():
            candidate = folder / f"{slugify(title)}-{counter}.{format}"
            counter += 1
        return candidate


def build_document_tools(studio: DocumentStudio) -> list[Tool]:
    return [
        Tool(
            name="create_document",
            description=(
                "Create a real document file and deliver it to the user in the chat (preview, Open, Download), "
                "like the Claude app. Formats: pdf, docx (Word), html, md, txt, csv. For pdf/docx/html/md, write "
                "`content` in Markdown (headings, lists, **bold**, tables); for txt/csv, the raw file content. "
                "Leave `path` empty unless the user named where to save it."
            ),
            parameters=object_schema(
                {
                    "title": {"type": "string", "description": "Document title, also used for the file name."},
                    "format": {"type": "string", "enum": list(FORMATS)},
                    "content": {"type": "string"},
                    "path": {"type": "string", "description": "Only if the user asked for a specific location."},
                },
                required=["title", "format", "content"],
            ),
            run=studio.create_document,
            confirmation_prompt=studio.create_document_confirmation,
        ),
        Tool(
            name="share_file",
            description=(
                "Show an existing file to the user in the chat with a preview, Open and Download buttons "
                "(e.g. a file you produced with run_command, or one the user asks to see)."
            ),
            parameters=object_schema(
                {"path": {"type": "string"}, "title": {"type": "string", "description": "Optional display title."}},
                required=["path"],
            ),
            run=studio.share_file,
        ),
    ]
