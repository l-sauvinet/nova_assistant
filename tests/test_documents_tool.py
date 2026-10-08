import io
from pathlib import Path

import docx
import pytest

from nova.core.messages import ToolCall
from nova.device.opener import open_command
from nova.tools.base import ToolError
from nova.tools.documents import DocumentStudio, build_document_tools
from nova.tools.file_access import FileAccessGuard
from nova.tools.registry import CANCELLED_BY_USER, ToolRegistry
from fakes import ScriptedConfirmer

COURSE = """# Droit pénal

Cours de **Master** avec des accents : é è à ç œ.

## I. Principes

- Légalité
- Non-rétroactivité

| Type | Peine |
|------|-------|
| Contravention | Amende |
| Crime | Réclusion |
"""


def studio(tmp_path: Path, confirm: bool = True) -> tuple[DocumentStudio, ScriptedConfirmer]:
    confirmer = ScriptedConfirmer(confirm)
    return DocumentStudio(FileAccessGuard([tmp_path], confirmer)), confirmer


def test_pdf_is_a_real_pdf_in_nova_folder(tmp_path: Path):
    output = studio(tmp_path)[0].create_document("Cours de droit pénal", "pdf", COURSE)
    media = output.media[0]
    path = Path(media.path)
    assert path == tmp_path.resolve() / "NOVA" / "cours-de-droit-penal.pdf"
    assert path.read_bytes().startswith(b"%PDF-") and path.stat().st_size > 1000
    assert media.kind == "file" and media.title == "Cours de droit pénal"


def test_docx_keeps_headings_lists_bold_and_tables(tmp_path: Path):
    path = Path(studio(tmp_path)[0].create_document("Cours", "docx", COURSE).media[0].path)
    document = docx.Document(io.BytesIO(path.read_bytes()))
    paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}
    assert paragraphs["Droit pénal"].style.name == "Heading 1"
    assert paragraphs["Légalité"].style.name == "List Bullet"
    assert any(run.bold and run.text == "Master" for run in paragraphs["Cours de Master avec des accents : é è à ç œ."].runs)
    table = document.tables[0]
    assert table.cell(0, 0).text == "Type" and table.cell(2, 1).text == "Réclusion"


def test_html_and_text_formats(tmp_path: Path):
    creator = studio(tmp_path)[0]
    html = Path(creator.create_document("Page", "html", "# Titre").media[0].path)
    assert "<h1>Titre</h1>" in html.read_text() and creator.create_document("x", "html", "a").media[0].kind == "html"
    csv = Path(creator.create_document("Budget", "csv", "mois;montant\njanvier;10").media[0].path)
    assert csv.read_text() == "mois;montant\njanvier;10"


def test_same_title_never_overwrites(tmp_path: Path):
    creator = studio(tmp_path)[0]
    first = creator.create_document("Lettre", "txt", "1").media[0].path
    second = creator.create_document("Lettre", "txt", "2").media[0].path
    assert first != second and Path(first).read_text() == "1"


def test_explicit_path_is_used_and_extension_added(tmp_path: Path):
    path = Path(studio(tmp_path)[0].create_document("CV", "pdf", "# CV", path="docs/mon-cv").media[0].path)
    assert path == tmp_path.resolve() / "docs" / "mon-cv.pdf"


def test_overwriting_an_explicit_path_asks_first(tmp_path: Path):
    (tmp_path / "cv.txt").write_text("ancien")
    creator, confirmer = studio(tmp_path, confirm=False)
    registry = ToolRegistry(build_document_tools(creator), confirmer)
    result = registry.execute(ToolCall(id="1", name="create_document", arguments={"title": "CV", "format": "txt", "content": "nouveau", "path": "cv.txt"}))
    assert result.content == CANCELLED_BY_USER and (tmp_path / "cv.txt").read_text() == "ancien"


def test_unknown_format_is_refused(tmp_path: Path):
    with pytest.raises(ToolError):
        studio(tmp_path)[0].create_document("x", "exe", "...")


def test_share_existing_file(tmp_path: Path):
    (tmp_path / "photo.png").write_bytes(b"png")
    (tmp_path / "rapport.pdf").write_bytes(b"%PDF")
    creator = studio(tmp_path)[0]
    assert creator.share_file("photo.png").media[0].kind == "image"
    shared = creator.share_file("rapport.pdf", title="Rapport annuel").media[0]
    assert shared.kind == "file" and shared.title == "Rapport annuel"
    with pytest.raises(ToolError):
        creator.share_file("absent.pdf")


def test_open_commands_per_platform(monkeypatch):
    assert open_command(Path("/w/a.pdf"), system="Windows") is None
    assert open_command(Path("/w/a.pdf"), system="Linux", wsl=False) == ["xdg-open", "/w/a.pdf"]
    monkeypatch.setattr("shutil.which", lambda name: "/usr/bin/wslview" if name == "wslview" else None)
    assert open_command(Path("/w/a.pdf"), system="Linux", wsl=True) == ["wslview", "/w/a.pdf"]
